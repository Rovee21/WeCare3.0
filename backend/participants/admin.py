from django import forms
from django.template.response import TemplateResponse
from django.urls import path
from django.contrib import admin, messages
from django.conf import settings
from django.core.mail import send_mail
from django.utils.html import format_html
from django.db.models import Count, Max, Q
from .models import Participant, CohortStartDate
import csv
import io
import re
import openpyxl
from django.http import HttpResponse

# Synonym lists used to pre-select the most likely dropdown option on the CSV/Excel
# import mapping screen. Matching is done against a normalized (lowercased,
# punctuation/whitespace-stripped) form of both the uploaded header and each synonym,
# so "E-mail", "e_mail", and "Email Address" all resolve to the same field.
# NOTE: 'group1' synonyms like "Study Cohort" are intentionally kept separate from
# the 'cohort' field's synonyms (recruitment wave) — these are different concepts.
FIELD_SYNONYMS = {
    'email': ['email', 'email address', 'e-mail', 'emailaddress', 'q_email'],
    'first_name': ['first name', 'firstname', 'first', 'given name'],
    'last_name': ['last name', 'lastname', 'last', 'surname', 'family name'],
    'gender': ['gender', 'sex'],
    'age': ['age', 'participant age'],
    'relationship': ['relationship', 'care relationship', 'relationship to care recipient', 'adrd relationship'],
    'group1': ['group1', 'group 1', 'study cohort', 'intervention group', 'cohort assignment'],
    'group2': ['group2', 'group 2', 'condition group', 'severity', 'adrd stage', 'stage', 'disease stage'],
    'group3': ['group3', 'group 3', 'stress group'],
    'cohort': ['cohort', 'recruitment cohort', 'recruitment wave', 'cohort number'],
}


def _normalize_header(value):
    return re.sub(r'[^a-z0-9]', '', (value or '').lower())


_NORMALIZED_FIELD_SYNONYMS = {
    field: {_normalize_header(s) for s in synonyms} | {_normalize_header(field)}
    for field, synonyms in FIELD_SYNONYMS.items()
}


def compute_auto_mapping(headers):
    """For each internal field key, return the first uploaded header whose normalized
    form matches one of its synonyms, or '' if nothing matches."""
    guess = {}
    for field_key, synonym_set in _NORMALIZED_FIELD_SYNONYMS.items():
        match = ''
        for h in headers:
            if _normalize_header(h) in synonym_set:
                match = h
                break
        guess[field_key] = match
    return guess


def parse_uploaded_participants_file(uploaded_file):
    """Parse a .csv or .xlsx upload into (headers, rows) where rows is a list of dicts
    keyed by header — the same shape csv.DictReader produces, regardless of format."""
    name = (uploaded_file.name or '').lower()
    if name.endswith('.xlsx'):
        wb = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [str(h).strip() if h is not None else '' for h in next(rows_iter)]
        except StopIteration:
            headers = []
        rows = []
        for raw_row in rows_iter:
            if raw_row is None or all(v is None for v in raw_row):
                continue
            row_dict = {}
            for i, header in enumerate(headers):
                if not header:
                    continue
                val = raw_row[i] if i < len(raw_row) else None
                if isinstance(val, float) and val.is_integer():
                    val = int(val)  # e.g. Age/Cohort read as 45.0 -> "45", not "45.0"
                row_dict[header] = '' if val is None else str(val)
            rows.append(row_dict)
        return headers, rows
    else:
        decoded = uploaded_file.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(decoded))
        rows = list(reader)
        headers = reader.fieldnames or []
        return headers, rows


class VoiceJournalInline(admin.TabularInline):
    model = None  # set below
    extra = 0
    can_delete = False
    fields = ["week_number", "duration_display", "vj_stress_level", "emotion_label", "transcription_status", "submitted_at"]
    readonly_fields = ["week_number", "duration_display", "vj_stress_level", "emotion_label", "transcription_status", "submitted_at"]
    verbose_name = "Voice Journal Submission"
    verbose_name_plural = "Voice Journal Submissions"
    ordering = ["week_number"]

    def duration_display(self, obj):
        m, s = divmod(obj.recording_seconds, 60)
        return f"{m}m {s:02d}s"
    duration_display.short_description = "Duration"

    def has_add_permission(self, request, obj=None):
        return False


class EngagementLogInline(admin.TabularInline):
    model = None  # set in _setup_inlines()
    extra = 0
    can_delete = False
    fields = [
        "week_number", "course_title",
        "total_time_display",
        "video_time_display", "video_watch_display", "text_time_display",
        "video_open_count", "emoji_taps_display", "logged_at",
    ]
    readonly_fields = [
        "week_number", "course_title",
        "total_time_display",
        "video_time_display", "video_watch_display", "text_time_display",
        "video_open_count", "emoji_taps_display", "logged_at",
    ]
    ordering = ["-logged_at"]
    verbose_name = "Engagement Event"
    verbose_name_plural = "Engagement Log"

    def _fmt(self, seconds):
        if not seconds:
            return "—"
        m, s = divmod(seconds, 60)
        return f"{m}:{s:02d}"

    def total_time_display(self, obj):
        total = obj.video_time_seconds + obj.text_time_seconds
        return self._fmt(total)
    total_time_display.short_description = "Total Time"

    def video_time_display(self, obj):
        return self._fmt(obj.video_time_seconds)
    video_time_display.short_description = "Video Time"

    def video_watch_display(self, obj):
        return self._fmt(obj.video_watch_seconds)
    video_watch_display.short_description = "Video Watch Time"

    def text_time_display(self, obj):
        return self._fmt(obj.text_time_seconds)
    text_time_display.short_description = "Text Time"

    def emoji_taps_display(self, obj):
        return obj.interactive_feature_count or "—"
    emoji_taps_display.short_description = "Emoji Taps"

    def has_add_permission(self, request, obj=None):
        return False


class SessionCompletionInline(admin.TabularInline):
    model = None  # set below
    extra = 0
    can_delete = False
    fields = ["session", "is_read", "started_at", "read_at"]
    readonly_fields = ["session", "is_read", "started_at", "read_at"]
    verbose_name = "Session Completion"
    verbose_name_plural = "Session Completions"

    def has_add_permission(self, request, obj=None):
        return False


class SessionOverrideInline(admin.TabularInline):
    model = None  # set below
    extra = 0
    fields = ["session", "override_type", "set_by", "created_at"]
    readonly_fields = ["set_by", "created_at"]
    verbose_name = "Session Override"
    verbose_name_plural = "Session Overrides"

    def formfield_for_foreignkey(self, db_field, request, **kwargs):
        if db_field.name == "session":
            from content.models import Session
            kwargs["queryset"] = Session.objects.filter(is_active=True).order_by("week_number", "day_number")
        return super().formfield_for_foreignkey(db_field, request, **kwargs)


def _setup_inlines():
    from journal.models import VoiceJournalEntry
    from content.models import ParticipantSession, EngagementLog, SessionOverride
    VoiceJournalInline.model = VoiceJournalEntry
    SessionCompletionInline.model = ParticipantSession
    EngagementLogInline.model = EngagementLog
    SessionOverrideInline.model = SessionOverride


_setup_inlines()


class ParticipantAdminForm(forms.ModelForm):
    class Meta:
        model = Participant
        fields = "__all__"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        # Cohort used to be a bare number spinner — replaced with a dropdown showing
        # each configured cohort's number + label, so an admin can pick the right one
        # by name instead of having to remember/guess numbers. Cohorts must be created
        # on the Cohort Start Dates page first; this dropdown only offers ones that
        # already exist there (plus the participant's own current value, even if it has
        # no CohortStartDate yet — e.g. a stray number from a CSV import — so opening an
        # existing participant never breaks) and an explicit "not yet assigned" option.
        known = {c.cohort: str(c) for c in CohortStartDate.objects.order_by("cohort")}
        current = getattr(self.instance, "cohort", None)
        if current is not None and current not in known:
            known[current] = f"Cohort {current} (no start date configured yet)"
        choices = [("", "— Not yet assigned —")] + sorted(known.items())
        self.fields["cohort"] = forms.TypedChoiceField(
            choices=choices, coerce=int, required=False, empty_value=None,
            label="Cohort",
            help_text="Which recruitment wave this participant belongs to. Add new cohorts "
                       "on the Cohort Start Dates page first, then pick one here.",
        )

    def clean(self):
        cleaned_data = super().clean()
        # Once a participant's cohort has started, clearing a baseline field they'd
        # already had set is a destructive regression during an active study — block
        # it outright. (Still-blank fields that were never filled in are fine — that
        # case is only warned about, in ParticipantAdmin.save_model, not blocked here.)
        if self.instance.pk and not self.instance.is_waitlisted():
            for field, label in Participant.BASELINE_FIELDS:
                old_value = getattr(self.instance, field)
                new_value = cleaned_data.get(field)
                if old_value and not new_value:
                    raise forms.ValidationError(
                        f'Cannot clear "{label}" — this participant\'s cohort has already '
                        "started. Editing baseline details back to blank isn't allowed "
                        "once the study is underway for them."
                    )
        return cleaned_data


@admin.register(Participant)
class ParticipantAdmin(admin.ModelAdmin):
    form = ParticipantAdminForm
    list_display = [
        "participant_id_display", "email", "language",
        "cohort", "cohort_label_display", "group1", "group2", "group3",
        "adrd_relationship_group",
        "is_enrolled", "enrollment_code_display", "current_week_display",
        "sessions_completed", "vj_count", "last_active",
        "code_status",
    ]
    list_filter = ["cohort"]
    search_fields = ["email"]
    readonly_fields = [
        "participant_id_display", "is_enrolled", "enrolled_at",
        "current_week_display", "latest_vj_stress",
        "user", "created_at", "updated_at",
    ]
    inlines = [SessionOverrideInline]
    list_per_page = 50

    def cohort_label_display(self, obj):
        c = CohortStartDate.objects.filter(cohort=obj.cohort).first()
        return (c.label if c and c.label else "—")
    cohort_label_display.short_description = "Cohort Label"
    cohort_label_display.admin_order_field = "cohort"

    def save_model(self, request, obj, form, change):
        super().save_model(request, obj, form, change)
        if not obj.is_waitlisted():
            missing = obj.missing_baseline_fields()
            if missing:
                self.message_user(
                    request,
                    f'"{obj}" — this participant\'s cohort has already started, but the '
                    f'following details are still missing: {", ".join(missing)}. Please '
                    "fill them in.",
                    level=messages.WARNING,
                )

    fieldsets = (
        ("Identity", {
            "fields": ("participant_id_display", "email", "first_name", "last_name", "language", "is_enrolled"),
        }),
        ("Demographics", {
            "fields": ("gender", "age"),
            "description": "From baseline survey.",
        }),
        ("Study Progress", {
            "fields": ("enrolled_at", "current_week_display", "enrollment_week"),
            "description": (
                "Content unlocks automatically over time (7 days/week, plus all of a "
                "week's sessions must be read to advance). To manually push a participant "
                "ahead of that schedule, raise Enrollment Week — it only ever moves them "
                "forward, never locks them out of progress they've already made."
            ),
        }),
        ("Caregiver Profile", {
            "fields": ("adrd_relationship_group", "group3"),
            "description": "Relationship and burden level from baseline survey.",
        }),
        ("Cohort Assignment", {
            "fields": ("group1", "group2"),
            "description": "Set from offline baseline survey data. Do not change after enrollment.",
        }),
        ("Recruitment Cohort", {
            "fields": ("cohort",),
            "description": (
                "Which recruitment wave this participant was enrolled under (Cohort 1, 2, ...). "
                "Unrelated to the clinical Study/Condition/Stress groups above."
            ),
        }),
        ("Clinical Indicators (from Voice Journal)", {
            "fields": ("latest_vj_stress",),
        }),
        ("Enrollment", {
            "fields": ("enrollment_code",),
            "classes": ("collapse",),
        }),
        ("System", {
            "fields": ("user", "created_at", "updated_at"),
            "classes": ("collapse",),
        }),
    )

    actions = [
        "generate_code_only",
        "generate_and_email_code",
        "generate_device_transfer_code",
        "send_notification_to_selected",
    ]

    def get_urls(self):
        urls = super().get_urls()
        custom = [
            path('<int:participant_id>/stats/', self.admin_site.admin_view(self.stats_view), name='participant_stats'),
            path('import-csv/', self.admin_site.admin_view(self.csv_import_view), name='participant_csv_import'),
            path('send-notification/', self.admin_site.admin_view(self.send_notification_view), name='participant_send_notification'),
            path('recordings/', self.admin_site.admin_view(self.recordings_view), name='participant_recordings'),
            path('download-import-template/', self.admin_site.admin_view(self.download_import_template_view), name='participant_download_import_template'),
        ]
        return custom + urls

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context['import_csv_url'] = '/admin/participants/participant/import-csv/'
        extra_context['send_notification_url'] = '/admin/participants/participant/send-notification/'
        extra_context['recordings_url'] = '/admin/participants/participant/recordings/'
        extra_context['download_template_url'] = '/admin/participants/participant/download-import-template/'
        return super().changelist_view(request, extra_context=extra_context)

    def download_import_template_view(self, request):
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Participants"

        headers = [
            "Email (Mandatory)", "First Name", "Last Name",
            "Gender", "Age", "Relationship", "Group1", "Group2", "Group3", "Cohort",
        ]
        example_row = [
            "jane.doe@example.com", "Jane", "Doe",
            "female", 68, "spouse", "intervention", "moderate", "high", 1,
        ]
        ws.append(headers)
        ws.append(example_row)

        email_cell = ws.cell(row=1, column=1)
        email_cell.font = Font(bold=True)
        email_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 22

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="participant_import_template.xlsx"'
        wb.save(response)
        return response

    def recordings_view(self, request):
        from journal.models import VoiceJournalEntry
        from journal.services import audio_download_url_for_entry, audio_filename_for_entry

        entries = VoiceJournalEntry.objects.select_related("participant").order_by("-submitted_at")
        rows = []
        for entry in entries:
            try:
                download_url = audio_download_url_for_entry(entry)
            except RuntimeError:
                download_url = None
            rows.append({
                "id": entry.id,
                "participant": entry.participant,
                "week_number": entry.week_number,
                "recording_seconds": entry.recording_seconds,
                "emotion": entry.get_emotion_label_display() or "—",
                "vj_stress_level": entry.vj_stress_level,
                "filename": audio_filename_for_entry(entry),
                "download_url": download_url,
                "submitted_at": entry.submitted_at,
            })

        context = {
            **self.admin_site.each_context(request),
            'title': 'Voice Recordings',
            'rows': rows,
        }
        return TemplateResponse(request, 'admin/participants/recordings_dashboard.html', context)
    
    def csv_import_view(self, request):
        from django.template.response import TemplateResponse

        # Step 3: Final import (no file, has 'import' flag)
        if request.method == 'POST' and request.POST.get('import') == '1':
            print("=== IMPORT TRIGGERED ===")
            csv_data = request.session.get('csv_import_data', '')
            print("Session data length:", len(csv_data))
            reader = csv.DictReader(io.StringIO(csv_data))
            mapping = {
                'email':                   request.POST.get('map_email', ''),
                'first_name':              request.POST.get('map_first_name', ''),
                'last_name':               request.POST.get('map_last_name', ''),
                'gender':                  request.POST.get('map_gender', ''),
                'age':                     request.POST.get('map_age', ''),
                'adrd_relationship_group': request.POST.get('map_relationship', ''),
                'group1':                  request.POST.get('map_group1', ''),
                'group2':                  request.POST.get('map_group2', ''),
                'group3':                  request.POST.get('map_group3', ''),
                'cohort':                  request.POST.get('map_cohort', ''),
            }

            valid_choices = {
                'group1': dict(Participant.GROUP1_CHOICES),
                'group2': dict(Participant.GROUP2_CHOICES),
                'group3': dict(Participant.GROUP3_CHOICES),
                'adrd_relationship_group': dict(Participant.RELATIONSHIP_CHOICES),
            }
            known_cohorts = set(CohortStartDate.objects.values_list('cohort', flat=True))

            def _clean_choice(value, field, row_email, warnings_list):
                # An unrecognized value (e.g. a typo like "superb" instead of
                # mild/moderate/severe) is left blank rather than stored as garbage —
                # a garbage value would otherwise look "complete" to
                # missing_baseline_fields() while matching no real targeting rule.
                if value and value not in valid_choices[field]:
                    warnings_list.append(
                        f'Row {row_email}: "{value}" is not a valid value for {field} — left blank.'
                    )
                    return ''
                return value

            created, skipped, errors, warnings = 0, 0, [], []
            new_participants = []
            for row in reader:
                email = row.get(mapping['email'], '').strip()
                if not email:
                    skipped += 1
                    continue
                # Guards against importing the Download Template's own example row if
                # someone forgets to delete it before uploading.
                if email.lower() == 'jane.doe@example.com':
                    skipped += 1
                    continue
                if Participant.objects.filter(email=email).exists():
                    skipped += 1
                    continue
                try:
                    age_val = row.get(mapping['age'], '').strip()
                    cohort_val = row.get(mapping['cohort'], '').strip()
                    cohort_num = int(cohort_val) if cohort_val.isdigit() else None
                    if cohort_num is not None and cohort_num not in known_cohorts:
                        warnings.append(
                            f'Row {email}: cohort {cohort_num} doesn\'t exist yet — left '
                            f"unassigned. Please create Cohort {cohort_num} on the Cohort "
                            "Start Dates page first, then edit this participant to assign it."
                        )
                        cohort_num = None
                    p = Participant(
                        email=email,
                        first_name=row.get(mapping['first_name'], '').strip(),
                        last_name=row.get(mapping['last_name'], '').strip(),
                        gender=row.get(mapping['gender'], '').strip().lower()[:10],
                        age=int(age_val) if age_val.isdigit() else None,
                        # Left blank (not guessed/defaulted) when not provided — same as
                        # a manually-added participant, so incomplete-profile warnings
                        # below are meaningful instead of being masked by a fabricated value.
                        adrd_relationship_group=_clean_choice(
                            row.get(mapping['adrd_relationship_group'], '').strip().lower()[:20],
                            'adrd_relationship_group', email, warnings,
                        ),
                        group1=_clean_choice(row.get(mapping['group1'], '').strip().lower()[:20], 'group1', email, warnings),
                        group2=_clean_choice(row.get(mapping['group2'], '').strip().lower()[:20], 'group2', email, warnings),
                        group3=_clean_choice(row.get(mapping['group3'], '').strip().lower()[:20], 'group3', email, warnings),
                        cohort=cohort_num,
                        language='en',
                        enrollment_week=1,
                    )
                    p.save()
                    p.generate_enrollment_code()
                    created += 1
                    new_participants.append(p)
                except Exception as e:
                    errors.append(f"Row {email}: {e}")

            self.message_user(request, f"Import complete: {created} created, {skipped} skipped, {len(errors)} errors.")
            if errors:
                for err in errors[:5]:
                    self.message_user(request, err, level='warning')
            if warnings:
                for warn in warnings[:10]:
                    self.message_user(request, warn, level=messages.WARNING)

            # Same "cohort already started but details incomplete" warning a manually-
            # added participant gets (ParticipantAdmin.save_model) — CSV import creates
            # participants directly, bypassing that path, so it needs its own check.
            incomplete = [
                (p, p.missing_baseline_fields()) for p in new_participants
                if not p.is_waitlisted() and p.missing_baseline_fields()
            ]
            for p, missing in incomplete[:10]:
                self.message_user(
                    request,
                    f'"{p}" — cohort already started but these details are still missing: '
                    f'{", ".join(missing)}. Please fill them in.',
                    level=messages.WARNING,
                )
            return self._redirect_to_changelist(request)

        # Step 2: Preview (has file upload)
        if request.method == 'POST' and 'csv_file' in request.FILES:
            uploaded_file = request.FILES['csv_file']
            headers, rows = parse_uploaded_participants_file(uploaded_file)

            # Re-serialize to canonical CSV text regardless of source format, so Step 3
            # can keep reading from session via the exact same csv.DictReader path.
            buf = io.StringIO()
            writer = csv.DictWriter(buf, fieldnames=headers)
            writer.writeheader()
            for row in rows:
                writer.writerow({h: row.get(h, '') for h in headers})
            request.session['csv_import_data'] = buf.getvalue()

            best_guess = compute_auto_mapping(headers)

            context = {
                **self.admin_site.each_context(request),
                'title': 'Import Participants from CSV',
                'headers': headers,
                'rows': rows[:5],
                'all_rows': rows,
                'best_guess': best_guess,
                'field_map': {
                    'email': 'Email *',
                    'first_name': 'First Name',
                    'last_name': 'Last Name',
                    'gender': 'Gender',
                    'age': 'Age',
                    'relationship': 'Care Relationship/ADRD relationship group - Spouse, Children/Adult Child, Other Relative',
                    'group1': 'Group1 - Study Cohort (Intervention/Control)',
                    'group2': 'Group2 - ADRD Stage (Mild/Moderate/Severe)',
                    'group3': 'Group3 - Stress Group (High/Low Stress)',
                    'cohort': 'Cohort / Recruitment Wave',
                },
            }
            return TemplateResponse(request, 'admin/participants/csv_preview.html', context)

        # Step 1: Show upload form
        context = {
            **self.admin_site.each_context(request),
            'title': 'Import Participants from CSV',
        }
        return TemplateResponse(request, 'admin/participants/csv_upload.html', context)

    def _redirect_to_changelist(self, request):
        from django.shortcuts import redirect
        return redirect('/admin/participants/participant/')

    def stats_view(self, request, participant_id):
        from content.models import EngagementLog, ParticipantSession
        from journal.models import VoiceJournalEntry
        from django.db.models import Sum

        participant = Participant.objects.get(pk=participant_id)

        raw_logs = EngagementLog.objects.filter(participant=participant).select_related('session').order_by('week_number', 'session__day_number')

        totals = raw_logs.aggregate(
            total_video=Sum('video_time_seconds'),
            total_video_watch=Sum('video_watch_seconds'),
            total_text=Sum('text_time_seconds'),
        )

        engagement_logs = []
        for log in raw_logs:
            total = log.video_time_seconds + log.text_time_seconds
            m, s = divmod(total, 60)
            log.total_time_display = f"{m}:{s:02d}" if total else "—"
            m, s = divmod(log.video_time_seconds, 60)
            log.video_time_display = f"{m}:{s:02d}" if log.video_time_seconds else "—"
            m, s = divmod(log.video_watch_seconds, 60)
            log.video_watch_display = f"{m}:{s:02d}" if log.video_watch_seconds else "—"
            m, s = divmod(log.text_time_seconds, 60)
            log.text_time_display = f"{m}:{s:02d}" if log.text_time_seconds else "—"
            engagement_logs.append(log)
        session_completions = ParticipantSession.objects.filter(participant=participant).select_related('session').order_by('session__week_number', 'session__day_number')
        vj_entries = VoiceJournalEntry.objects.filter(participant=participant).order_by('week_number')
        latest_vj_entry = VoiceJournalEntry.objects.filter(participant=participant).order_by('-submitted_at').first()

        def fmt(seconds):
            if not seconds:
                return '—'
            m, s = divmod(seconds, 60)
            return f'{m}:{s:02d}'

        total_seconds = (totals['total_video'] or 0) + (totals['total_text'] or 0)

        context = {
            **self.admin_site.each_context(request),
            'participant': participant,
            'engagement_logs': engagement_logs,
            'session_completions': session_completions,
            'vj_entries': vj_entries,
            'summary': {
                'sessions_read': session_completions.filter(is_read=True).count(),
                'total_time': fmt(total_seconds),
                'total_video_time': fmt(totals['total_video'] or 0),
                'total_video_watch_time': fmt(totals['total_video_watch'] or 0),
                'total_text_time': fmt(totals['total_text'] or 0),
                'vj_submitted': vj_entries.count(),
                'latest_stress': latest_vj_entry.vj_stress_level if latest_vj_entry else '—',
                'latest_emotion': latest_vj_entry.get_emotion_label_display() if latest_vj_entry and latest_vj_entry.emotion_label else '—',
                'last_active': raw_logs.order_by('-logged_at').first().logged_at if raw_logs.exists() else None,
            },
            'fmt': fmt,
            'title': f'{participant.participant_id} — Stats',
        }
        return TemplateResponse(request, 'admin/participant_stats.html', context)

    def _send_notification_context(self, request, selected_participants, participant_ids_param, form_values=None):
        if form_values is None:
            form_values = {
                'target_mode': 'selected' if selected_participants else 'single',
                'timing': 'now',
                'title': '',
                'body': '',
                'single_participant_id': '',
                'cohort_group1': '',
                'cohort_group2': '',
                'cohort_group3': '',
                'scheduled_for': '',
            }
        return {
            **self.admin_site.each_context(request),
            'title': 'Send Notification',
            'participants': Participant.objects.all().order_by('email'),
            'group1_choices': Participant.GROUP1_CHOICES,
            'group2_choices': Participant.GROUP2_CHOICES,
            'group3_choices': Participant.GROUP3_CHOICES,
            'selected_participants': selected_participants,
            'participant_ids_param': participant_ids_param,
            'form_values': form_values,
        }

    def send_notification_view(self, request):
        from django.shortcuts import redirect
        from django.utils import timezone as dj_timezone
        from django.utils.dateparse import parse_datetime
        from content.models import NotificationLog
        from participants.notifications import (
            send_and_log_notification,
            send_and_log_notification_bulk,
            schedule_notification,
            schedule_notification_bulk,
        )

        participant_ids_param = request.POST.get('participant_ids') or request.GET.get('participant_ids', '')
        selected_ids = [int(x) for x in participant_ids_param.split(',') if x.strip().isdigit()]
        selected_participants = list(Participant.objects.filter(pk__in=selected_ids)) if selected_ids else []

        if request.method == 'POST':
            target_mode = request.POST.get('target_mode', 'single')
            title = request.POST.get('title', '').strip()
            body = request.POST.get('body', '').strip()
            timing = request.POST.get('timing', 'now')
            scheduled_for_raw = request.POST.get('scheduled_for', '')

            def render_with_error(error):
                form_values = {
                    'target_mode': target_mode,
                    'timing': timing,
                    'title': title,
                    'body': body,
                    'single_participant_id': request.POST.get('single_participant_id', ''),
                    'cohort_group1': request.POST.get('cohort_group1', ''),
                    'cohort_group2': request.POST.get('cohort_group2', ''),
                    'cohort_group3': request.POST.get('cohort_group3', ''),
                    'scheduled_for': scheduled_for_raw,
                }
                context = self._send_notification_context(
                    request, selected_participants, participant_ids_param, form_values=form_values
                )
                context['error'] = error
                return TemplateResponse(request, 'admin/participants/send_notification.html', context)

            if not title or not body:
                return render_with_error("Title and body are both required.")

            if target_mode == 'single':
                pid = request.POST.get('single_participant_id')
                participants = list(Participant.objects.filter(pk=pid)) if pid else []
            elif target_mode == 'cohort':
                filters = {}
                for field in ('group1', 'group2', 'group3'):
                    value = request.POST.get(f'cohort_{field}', '')
                    if value:
                        filters[field] = value
                participants = list(Participant.objects.filter(**filters))
            elif target_mode == 'everyone':
                participants = list(Participant.objects.all())
            elif target_mode == 'selected':
                participants = selected_participants
            else:
                participants = []

            if not participants:
                return render_with_error("No participants matched the selected target — nothing was sent.")

            scheduled_dt = None
            if timing == 'schedule':
                scheduled_dt = parse_datetime(scheduled_for_raw)
                if scheduled_dt and dj_timezone.is_naive(scheduled_dt):
                    scheduled_dt = dj_timezone.make_aware(scheduled_dt)
                if not scheduled_dt:
                    return render_with_error("Please choose a valid date/time to schedule for.")

            if timing == 'schedule':
                if len(participants) == 1:
                    logs = [schedule_notification(participants[0], title, body, scheduled_dt)]
                else:
                    logs = schedule_notification_bulk(participants, title, body, scheduled_dt)
                pending = sum(1 for l in logs if l.status == NotificationLog.STATUS_PENDING)
                sent_now = len(logs) - pending
                when = dj_timezone.localtime(scheduled_dt).strftime('%b %d, %Y %I:%M %p')
                if pending == len(logs):
                    msg = f"Scheduled for {when} — will send to {len(logs)} participant(s) when processed."
                else:
                    msg = (
                        f"{pending} scheduled for {when}; {sent_now} sent immediately "
                        f"(their scheduled time had already passed)."
                    )
                self.message_user(request, msg)
            else:
                if len(participants) == 1:
                    log = send_and_log_notification(participants[0], title, body)
                    summary = {
                        'sent': int(log.status == NotificationLog.STATUS_SENT),
                        'skipped_no_token': int(log.status == NotificationLog.STATUS_NO_TOKEN),
                        'failed': int(log.status == NotificationLog.STATUS_FAILED),
                    }
                else:
                    summary = send_and_log_notification_bulk(participants, title, body)
                msg = f"Sent to {summary['sent']} participant(s)"
                extras = []
                if summary['skipped_no_token']:
                    extras.append(f"{summary['skipped_no_token']} skipped — no device token")
                if summary['failed']:
                    extras.append(f"{summary['failed']} failed")
                if extras:
                    msg += " (" + ", ".join(extras) + ")"
                msg += "."
                self.message_user(request, msg)

            return self._redirect_to_changelist(request)

        context = self._send_notification_context(request, selected_participants, participant_ids_param)
        return TemplateResponse(request, 'admin/participants/send_notification.html', context)

    @admin.action(description="Send Notification to selected")
    def send_notification_to_selected(self, request, queryset):
        from django.shortcuts import redirect
        ids = ",".join(str(pk) for pk in queryset.values_list("pk", flat=True))
        return redirect(f"/admin/participants/participant/send-notification/?participant_ids={ids}")

    def change_view(self, request, object_id, form_url='', extra_context=None):
        extra_context = extra_context or {}
        extra_context['show_history'] = False
        extra_context['stats_url'] = f'/admin/participants/participant/{object_id}/stats/'
        return super().change_view(request, object_id, form_url, extra_context)

    def save_formset(self, request, form, formset, change):
        instances = formset.save(commit=False)
        for obj in instances:
            if hasattr(obj, "set_by_id") and not obj.set_by_id:
                obj.set_by = request.user
            obj.save()
        formset.save_m2m()
        for obj in formset.deleted_objects:
            obj.delete()

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        return qs.annotate(
            _sessions_completed=Count("session_states", filter=Q(session_states__is_read=True), distinct=True),
            _vj_count=Count("journal_entries", distinct=True),
            _last_active=Max("engagement_logs__logged_at"),
        )

    def participant_id_display(self, obj):
        return obj.participant_id if obj.pk else "—"
    participant_id_display.short_description = "Participant ID"

    def current_week_display(self, obj):
        if not obj.is_enrolled:
            return "—"

        start = obj.program_start_date()
        if obj.is_waitlisted():
            if obj.cohort is None:
                detail = "No cohort has been assigned yet."
            elif start is None:
                detail = f"Cohort {obj.cohort} has no start date configured yet."
            else:
                detail = f"Cohort {obj.cohort} starts {start.strftime('%Y-%b-%d')} (in the future)."
            return format_html(
                '<span style="background:#fff3e0;color:#e65100;padding:2px 8px;border-radius:10px;font-weight:600;">Waitlisted</span>'
                '<div style="color:#888;font-size:12px;margin-top:4px;">{}</div>',
                detail,
            )

        week = obj.current_week_number
        automatic = obj.automatic_gated_week()
        detail = f"Cohort {obj.cohort} started {start.strftime('%Y-%b-%d')}." if start else ""
        if obj.enrollment_week and obj.enrollment_week > automatic:
            detail += (
                f" Manually advanced via Enrollment Week override — calendar/completion "
                f"alone would currently give Week {automatic}."
            )
        return format_html(
            '<span style="background:#e8f5e9;color:#2e7d32;padding:2px 8px;border-radius:10px;font-weight:600;">Week {}</span>'
            '<div style="color:#888;font-size:12px;margin-top:4px;">{}</div>',
            week, detail,
        )
    current_week_display.short_description = "Current Week"

    def latest_vj_stress(self, obj):
        entry = obj.journal_entries.order_by("-submitted_at").first()
        if not entry:
            return format_html('<span style="color:#999;">No submissions yet</span>')
        if entry.vj_stress_level is None:
            return format_html('<span style="color:#999;">— awaiting report</span>')
        level = entry.vj_stress_level
        if level <= 3:
            color = "#2e7d32"
        elif level <= 6:
            color = "#f57c00"
        else:
            color = "#c62828"
        return format_html(
            '<span style="color:{};font-weight:600;">{}/10</span> '
            '<span style="color:#999;font-size:11px;">(Week {})</span>',
            color, level, entry.week_number,
        )
    latest_vj_stress.short_description = "VJ Stress Level"

    def sessions_completed(self, obj):
        count = getattr(obj, "_sessions_completed", 0)
        color = "#2e7d32" if count > 0 else "#999"
        return format_html('<span style="color:{};">{}</span>', color, count)
    sessions_completed.short_description = "Sessions Read"
    sessions_completed.admin_order_field = "_sessions_completed"

    def vj_count(self, obj):
        count = getattr(obj, "_vj_count", 0)
        color = "#1565c0" if count > 0 else "#999"
        return format_html('<span style="color:{};">{} / 7</span>', color, count)
    vj_count.short_description = "VJ Submitted"
    vj_count.admin_order_field = "_vj_count"

    def last_active(self, obj):
        dt = getattr(obj, "_last_active", None)
        if not dt:
            return format_html('<span style="color:#999;">Never</span>')
        return format_html('<span title="{}">{}</span>', dt.strftime("%Y-%m-%d %H:%M"), dt.strftime("%b %d"))
    last_active.short_description = "Last Active"
    last_active.admin_order_field = "_last_active"

    def enrollment_code_display(self, obj):
        if not obj.enrollment_code:
            return format_html('<span style="color:#999;">—</span>')
        return format_html(
            '<code style="background:#fff8e1;padding:2px 8px;border-radius:3px;letter-spacing:1px;">{}</code>',
            obj.enrollment_code,
        )
    enrollment_code_display.short_description = "Enrollment Code"
    enrollment_code_display.admin_order_field = "enrollment_code"

    def code_status(self, obj):
        if obj.is_enrolled:
            return format_html('<span style="color:#2e7d32;font-weight:600;">✓ Enrolled</span>')
        if obj.enrollment_code:
            return format_html(
                '<code style="background:#fff8e1;padding:2px 8px;border-radius:3px;letter-spacing:1px;">{}</code>',
                obj.enrollment_code,
            )
        return format_html('<span style="color:#e53935;">No code</span>')
    code_status.short_description = "Code / Status"

    @admin.action(description="Generate enrollment code (no email)")
    def generate_code_only(self, request, queryset):
        count = 0
        for p in queryset.filter(is_enrolled=False):
            p.generate_enrollment_code()
            count += 1
        self.message_user(request, f"Generated enrollment codes for {count} participant(s).")

    @admin.action(description="Generate code and email participant")
    def generate_and_email_code(self, request, queryset):
        sent, skipped = 0, 0
        for p in queryset:
            if p.is_enrolled:
                skipped += 1
                continue
            code = p.generate_enrollment_code()
            send_mail(
                subject="Your WeCare Enrollment Code",
                message=(
                    f"Hello {p.participant_id},\n\n"
                    f"Your WeCare program enrollment code is:\n\n"
                    f"    {code}\n\n"
                    f"Please enter this code in the WeCare app to get started.\n\n"
                    f"If you have any questions, reply to this email or contact us at wecaremason@gmail.com\n\n"
                    f"— The WeCare Research Team"
                ),
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[p.email],
                fail_silently=False,
            )
            sent += 1
        msg = f"Sent enrollment codes to {sent} participant(s)."
        if skipped:
            msg += f" Skipped {skipped} already-enrolled participant(s)."
        self.message_user(request, msg)

    @admin.action(description="Generate device transfer code")
    def generate_device_transfer_code(self, request, queryset):
        count = 0
        for p in queryset:
            p.generate_enrollment_code()
            count += 1
        self.message_user(request, f"Generated device transfer code(s) for {count} participant(s).")


class CohortStartDateAdminForm(forms.ModelForm):
    class Meta:
        model = CohortStartDate
        fields = "__all__"

    def clean(self):
        from django.utils import timezone as dj_timezone

        cleaned_data = super().clean()
        cohort = cleaned_data.get("cohort")
        start_date = cleaned_data.get("program_start_date")
        if cohort is not None and start_date and start_date <= dj_timezone.localdate():
            incomplete = []
            for p in Participant.objects.filter(cohort=cohort):
                missing = p.missing_baseline_fields()
                if missing:
                    incomplete.append(f"{p.participant_id} ({', '.join(missing)})")
            if incomplete:
                raise forms.ValidationError(
                    f"Cannot set Cohort {cohort}'s start date to today or the past — the "
                    f"following participants still have incomplete baseline details: "
                    f"{'; '.join(incomplete)}. Fill those in first."
                )
        return cleaned_data


@admin.register(CohortStartDate)
class CohortStartDateAdmin(admin.ModelAdmin):
    form = CohortStartDateAdminForm
    list_display = ["cohort", "label", "program_start_date_display", "updated_at"]

    def program_start_date_display(self, obj):
        return obj.program_start_date.strftime("%Y-%b-%d")
    program_start_date_display.short_description = "Program Start Date"
    program_start_date_display.admin_order_field = "program_start_date"